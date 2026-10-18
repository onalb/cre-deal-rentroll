from __future__ import annotations

import csv
import io
import re
from datetime import date

from openpyxl import load_workbook

from cre_rentroll.schemas import CanonicalRentRoll, UnitRow

PII_HEADERS = re.compile(
    r"^(tenant|resident|name|email|phone|ssn|ssn#|social|guarantor).*$",
    re.I,
)
CHARGE_TOTAL = re.compile(r"charge\s*total", re.I)
AMT = re.compile(
    r"scheduled\s*charges|^amount$|charge\s*amount|charge\s*total",
    re.I,
)
INPLACE = re.compile(
    r"inplace\s*rent|in[_\s-]*place\s*rent|current\s*rent|market\s*rent|^rent$|lease\s*rent",
    re.I,
)
OCC = re.compile(r"occ(upancy)?\s*status|unit\s*status|lease\s*status|(?<!sub)status$|^occupied$", re.I)
UNIT = re.compile(r"^(unit|unit\s*(id|no|number|#)|apt|bldg[\s-]*unit|building[\s-]*unit)$", re.I)
LEASE_END = re.compile(r"lease\s*end|end\s*date|move[\s-]*out", re.I)
LEASE_START = re.compile(r"lease\s*start|start\s*date|move[\s-]*in", re.I)
_HEADER_HINT = re.compile(
    r"unit|lease|status|rent|charge|occ|bldg|sqft|sq\s*ft|amount|type|mtm",
    re.I,
)
_REAL_UNIT = re.compile(r"^\d+\s*[-/]\s*\d+")
TYPE = re.compile(r"unit\s*type|floorplan|bed", re.I)
SQFT = re.compile(r"sq\s*ft|sqft|area", re.I)
MTM = re.compile(r"m(onth)?[\s-]*t(o)?[\s-]*m(onth)?|mtm", re.I)
ADDR = re.compile(r"address|street|location", re.I)
PROP = re.compile(r"property|community|asset|name$", re.I)
ASOF = re.compile(r"as[\s-]*of|reportdate|report\s*date", re.I)
VACANT_TOKENS = re.compile(r"vacant|empty|down", re.I)
_FILE_DATE = re.compile(r"(?<!\d)(\d{1,2})[-/_](\d{1,2})[-/_](\d{2}|\d{4})(?!\d)")


def _norm(h: str) -> str:
    return re.sub(r"\s+", " ", (h or "").strip())


def _is_pii(header: str) -> bool:
    h = _norm(header)
    if PROP.match(h) and "tenant" not in h.lower():
        return False
    return bool(PII_HEADERS.match(h))


def _num(val: object) -> float:
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(",", "").replace("$", "")
    if s in {"", "-", "None"}:
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _read_tabular(filename: str, raw: bytes) -> tuple[list[str], list[list[object]]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xls")):
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        rows = [r for r in rows if any(c is not None and str(c).strip() != "" for c in r)]
        if not rows:
            return [], []
        return _headers_and_body(rows)
    text = raw.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = list(reader)
    if not rows:
        return [], []
    return _headers_and_body(rows)


def _headers_and_body(rows: list[list[object]]) -> tuple[list[str], list[list[object]]]:
    hi = _find_header_row(rows)
    raw = rows[hi]
    headers = [_norm(str(c) if c is not None else "") for c in raw]
    body = rows[hi + 1 :]
    width = max(len(headers), max((len(r) for r in body), default=0))
    headers = (headers + [""] * width)[:width]
    return headers, body


def _find_header_row(rows: list[list[object]]) -> int:
    best_i, best = 0, -1
    for i, row in enumerate(rows[:45]):
        cells = [_norm(str(c or "")) for c in row]
        nonempty = [c for c in cells if c]
        if len(nonempty) < 3:
            continue
        score = sum(1 for c in nonempty if _HEADER_HINT.search(c))
        if score >= 3 and score > best:
            best_i, best = i, score
    return best_i


def detect_layout(headers: list[str], rows: list[list[object]]) -> str:
    joined = " ".join(headers).lower()
    if "charge" in joined or any(CHARGE_TOTAL.search(h) for h in headers):
        return "realpage"
    blob = " ".join(str(c) for row in rows[:40] for c in row)
    if CHARGE_TOTAL.search(blob):
        return "realpage"
    if any(INPLACE.search(h) for h in headers) and any(OCC.search(h) for h in headers):
        return "one_row"
    if any(INPLACE.search(h) for h in headers):
        return "one_row"
    return "one_row"


def _col(headers: list[str], pattern: re.Pattern[str]) -> int | None:
    for i, h in enumerate(headers):
        if pattern.search(h):
            return i
    return None


def _meta_from_grid(headers: list[str], rows: list[list[object]]) -> dict[str, str | None]:
    name = as_of = address = None
    # Header-row metadata (property name sitting in a column)
    pi = _col(headers, PROP)
    ai = _col(headers, ADDR)
    di = _col(headers, ASOF)
    if rows:
        if pi is not None and pi < len(rows[0]):
            name = str(rows[0][pi] or "").strip() or None
        if ai is not None and ai < len(rows[0]):
            address = str(rows[0][ai] or "").strip() or None
        if di is not None and di < len(rows[0]):
            as_of = str(rows[0][di] or "").strip() or None
    # Key/value first-column sheets
    for row in rows[:12]:
        if not row:
            continue
        key = _norm(str(row[0] or ""))
        val = _norm(str(row[1] or "")) if len(row) > 1 else ""
        if not val:
            continue
        if PROP.search(key) and not name:
            name = val
        if ASOF.search(key) and not as_of:
            as_of = val
        if ADDR.search(key) and not address:
            address = val
    return {"property_name": name, "as_of": as_of, "suggested_address": address}


def _geo_confidence(name: str | None, address: str | None) -> str:
    blob = f"{name or ''} {address or ''}".lower()
    if "chambers hill" in blob:
        return "low"
    if address and re.search(r"\d+", address) and re.search(r"[A-Za-z]{3,}", address):
        if "," in address or re.search(r"\b[A-Z]{2}\b", address or ""):
            return "high"
        return "medium"
    return "low"


def _vacant(occ: str, rent: float) -> bool:
    if VACANT_TOKENS.search(occ):
        return True
    if re.search(r"occup", occ, re.I):
        return False
    return rent <= 0 and not occ


def _keep_uid(uid: str, prefer_bldg: bool) -> bool:
    if not uid or uid.lower() in {"unit", "total", "property"}:
        return False
    if PROP.match(uid) or ASOF.match(uid):
        return False
    if re.search(
        r"status summary|^description$|^total (units|occupied)|occupied no notice|^unit details$",
        uid,
        re.I,
    ):
        return False
    if prefer_bldg:
        return bool(_REAL_UNIT.match(uid))
    return True


def _prefer_bldg(rows: list[list[object]], ui: int) -> bool:
    n = 0
    for row in rows:
        if ui < len(row) and _REAL_UNIT.match(str(row[ui] or "").strip()):
            n += 1
            if n >= 3:
                return True
    return False


def parse_realpage(headers: list[str], rows: list[list[object]]) -> list[UnitRow]:
    ui = _col(headers, UNIT) or 0
    ti = _col(headers, TYPE)
    si = _col(headers, SQFT)
    oi = _col(headers, OCC)
    mi = _col(headers, MTM)
    lsi = _col(headers, LEASE_START)
    lei = _col(headers, LEASE_END)
    prefer = _prefer_bldg(rows, ui)
    charge_i = _col(headers, re.compile(r"charge\s*code", re.I))
    if charge_i is None:
        charge_i = _col(headers, re.compile(r"charge\s*desc|description", re.I))
    amt_i = _col(headers, AMT)
    mri = _col(headers, INPLACE)
    if amt_i is None:
        amt_i = mri
    if amt_i is None:
        amt_i = len(headers) - 1 if headers else 0

    by_unit: dict[str, UnitRow] = {}
    for row in rows:
        if not row or ui >= len(row):
            continue
        uid = str(row[ui] or "").strip()
        if not _keep_uid(uid, prefer):
            continue
        charge = str(row[charge_i] or "").strip() if charge_i is not None and charge_i < len(row) else ""
        amt = _num(row[amt_i] if amt_i < len(row) else 0)
        unit = by_unit.get(uid)
        if unit is None:
            occ = str(row[oi] or "") if oi is not None and oi < len(row) else ""
            ls = str(row[lsi] or "").strip() if lsi is not None and lsi < len(row) else ""
            le = str(row[lei] or "").strip() if lei is not None and lei < len(row) else ""
            unit = UnitRow(
                unit_id=uid,
                unit_type=str(row[ti] or "").strip() or None if ti is not None and ti < len(row) else None,
                sqft=_num(row[si]) or None if si is not None and si < len(row) else None,
                occupancy=occ or None,
                mtm=bool(mi is not None and mi < len(row) and str(row[mi] or "").strip() not in {"", "0", "N", "No"}),
                lease_start=ls or None,
                lease_end=le or None,
            )
            by_unit[uid] = unit
            if mri is not None and mri < len(row) and unit.in_place_rent == 0:
                mkt = _num(row[mri])
                if mkt:
                    unit.charge_codes["_market_rent"] = mkt
        if charge:
            unit.charge_codes[charge] = amt
        if amt and (CHARGE_TOTAL.search(charge) or re.search(r"^rent(\s+charge)?$", charge, re.I)):
            unit.in_place_rent = amt
        elif amt and unit.in_place_rent == 0:
            unit.in_place_rent = amt
    for unit in by_unit.values():
        if unit.in_place_rent == 0 and unit.charge_codes:
            extras = {k: v for k, v in unit.charge_codes.items() if k != "_market_rent"}
            unit.in_place_rent = max(extras.values()) if extras else unit.charge_codes.get("_market_rent", 0.0)
        if unit.in_place_rent == 0:
            unit.in_place_rent = unit.charge_codes.pop("_market_rent", 0.0)
        else:
            unit.charge_codes.pop("_market_rent", None)
        unit.vacant = _vacant(unit.occupancy or "", unit.in_place_rent)
        if unit.vacant:
            unit.in_place_rent = 0.0
    return list(by_unit.values())


def parse_one_row(headers: list[str], rows: list[list[object]]) -> list[UnitRow]:
    ui = _col(headers, UNIT)
    ri = _col(headers, INPLACE)
    oi = _col(headers, OCC)
    ti = _col(headers, TYPE)
    si = _col(headers, SQFT)
    mi = _col(headers, MTM)
    lsi = _col(headers, LEASE_START)
    lei = _col(headers, LEASE_END)
    if ui is None:
        ui = 0
    prefer = _prefer_bldg(rows, ui)
    units: list[UnitRow] = []
    seen: set[str] = set()
    for row in rows:
        if not row or ui >= len(row):
            continue
        uid = str(row[ui] or "").strip()
        if not _keep_uid(uid, prefer) or uid in seen:
            continue
        rent = _num(row[ri]) if ri is not None and ri < len(row) else 0.0
        occ = str(row[oi] or "").strip() if oi is not None and oi < len(row) else ""
        ls = str(row[lsi] or "").strip() if lsi is not None and lsi < len(row) else ""
        le = str(row[lei] or "").strip() if lei is not None and lei < len(row) else ""
        vacant = _vacant(occ, rent)
        mtm = bool(mi is not None and mi < len(row) and str(row[mi] or "").lower() in {"y", "yes", "true", "1", "mtm"})
        if not mtm and le and re.search(r"mtm|month", le, re.I):
            mtm = True
        units.append(
            UnitRow(
                unit_id=uid,
                unit_type=str(row[ti] or "").strip() or None if ti is not None and ti < len(row) else None,
                sqft=_num(row[si]) or None if si is not None and si < len(row) else None,
                occupancy=occ or None,
                in_place_rent=0.0 if vacant else rent,
                vacant=vacant,
                mtm=mtm,
                lease_start=ls or None,
                lease_end=le or None,
            )
        )
        seen.add(uid)
    return units


def _year(yy: str) -> int:
    y = int(yy)
    if y < 100:
        return 2000 + y if y < 70 else 1900 + y
    return y


def normalize_as_of(raw: object | None) -> str | None:
    if raw is None:
        return None
    if isinstance(raw, date):
        return raw.isoformat()
    s = str(raw).strip()
    if not s or s.lower() in {"none", "unknown", "n/a"}:
        return None
    m = _FILE_DATE.search(s.replace(".", "-"))
    if not m:
        return s
    a, b, yy = m.group(1), m.group(2), m.group(3)
    year = _year(yy)
    for month, day in ((int(a), int(b)), (int(b), int(a))):
        try:
            return date(year, month, day).isoformat()
        except ValueError:
            continue
    return s


_ASOF_IN_MSG = re.compile(r"as(?:[\s-]*of)\s+(\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4})", re.I)


def as_of_from_message(message: str) -> str | None:
    m = _ASOF_IN_MSG.search(message or "")
    if not m:
        return None
    return normalize_as_of(m.group(1))


def as_of_from_filename(filename: str) -> str | None:
    stem = filename.rsplit(".", 1)[0]
    m = _FILE_DATE.search(stem.replace(".", "-"))
    if not m:
        return None
    return normalize_as_of(m.group(0))


def ingest_bytes(filename: str, raw: bytes) -> CanonicalRentRoll:
    headers, rows = _read_tabular(filename, raw)
    if not headers:
        raise ValueError("empty spreadsheet")
    # Drop PII columns before parse
    keep = [i for i, h in enumerate(headers) if h and not _is_pii(h)]
    headers = [headers[i] for i in keep]
    slim_rows: list[list[object]] = []
    for row in rows:
        slim_rows.append([row[i] if i < len(row) else None for i in keep])
    layout = detect_layout(headers, slim_rows)
    meta = _meta_from_grid(headers, slim_rows)
    units = parse_realpage(headers, slim_rows) if layout == "realpage" else parse_one_row(headers, slim_rows)
    if not units:
        raise ValueError("no unit rows parsed")
    name = meta["property_name"] or filename.rsplit(".", 1)[0]
    address = meta["suggested_address"]
    as_of = normalize_as_of(meta["as_of"]) or as_of_from_filename(filename)
    return CanonicalRentRoll(
        property_name=str(name) if name else None,
        as_of=as_of,
        suggested_address=str(address) if address else None,
        geo_confidence=_geo_confidence(str(name) if name else None, str(address) if address else None),
        layout=layout,  # type: ignore[arg-type]
        units=units,
        source_filename=filename,
    )


def ack_text(roll: CanonicalRentRoll) -> str:
    name = roll.property_name or "the rent roll"
    line = f"Saved {name} — {roll.unit_count} units"
    if roll.as_of:
        line += f", as of {roll.as_of}"
    return line + ". Tenant names were removed."
