"""Read occupied-unit features from the interview rent rolls. No resident names."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import openpyxl

FEATURE_NAMES = [
    "sqft",
    "beds",
    "reno",
    "retail",
]

_UNIT_RE = re.compile(r"^\d+-\d+")
_BEDS_RE = re.compile(r"(\d+)\s*x\s*\d+", re.I)


def months_between(start: datetime | None, end: datetime | None) -> float | None:
    if start is None or end is None:
        return None
    return (end.year - start.year) * 12 + (end.month - start.month) + (end.day - start.day) / 30.0


def beds_from_plan(plan: str) -> float | None:
    match = _BEDS_RE.search(plan or "")
    if not match:
        return None
    return float(match.group(1))


def feature_row(unit: dict) -> list[float]:
    """Physical facts only: size, bedrooms, renovation, retail."""
    beds = unit.get("beds")
    return [
        float(unit.get("sqft") or 0.0),
        float(beds) if beds is not None else 0.0,
        1.0 if unit.get("reno") else 0.0,
        1.0 if unit.get("retail") else 0.0,
    ]


def load_training_rows(cordoba: Path, bay_view: Path, chambers: Path) -> list[dict]:
    rows = []
    rows.extend(_cordoba(cordoba, datetime(2021, 4, 30)))
    rows.extend(_rediq(bay_view, "Bay View Gardens", datetime(2021, 5, 25)))
    rows.extend(_rediq(chambers, "Chambers Hill", datetime(2020, 11, 12)))
    return rows


def _rent_code(code: str) -> bool:
    cl = code.lower().strip()
    if "pet" in cl or "concession" in cl or "parking" in cl:
        return False
    return cl in {"rent charge", "rent", "base rent", "housing voucher rent"} or cl.startswith("rent charge")


def _cordoba(path: Path, asof: datetime) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = list(wb.active.iter_rows(values_only=True))
    wb.close()
    units: list[dict] = []
    cur: dict | None = None

    def apply(code: str, amt: float | None) -> None:
        if cur is None or not code or amt is None:
            return
        if "total" in code.lower() or code == "-":
            return
        cl = code.lower()
        if "month-to-month" in cl:
            cur["mtm"] = True
        elif _rent_code(code):
            cur["in_place"] += amt

    def flush() -> None:
        if cur is not None:
            units.append(cur)

    for row in sheet[7:]:
        uid = row[0]
        if isinstance(uid, str) and _UNIT_RE.match(uid.strip()):
            flush()
            status = str(row[3] or "").strip()
            low = status.lower()
            if "model" in low or "excluded" in low:
                occ = "excluded"
            elif "vacant" in low:
                occ = "vacant"
            elif "occup" in low or "notice" in low:
                occ = "occupied"
            else:
                occ = "other"
            plan = str(row[1] or "")
            cur = {
                "property": "La Cordoba",
                "as_of": asof.date().isoformat(),
                "plan": plan,
                "sqft": float(row[2]) if isinstance(row[2], (int, float)) else None,
                "beds": beds_from_plan(plan),
                "reno": "reno" in plan.lower(),
                "mtm": False,
                "in_place": 0.0,
                "market": float(row[5]) if isinstance(row[5], (int, float)) else None,
                "occ": occ,
                "notice": ("notice" in low and "no notice" not in low) or isinstance(row[13], datetime),
                "retail": False,
                "tenure_months": months_between(row[10] if isinstance(row[10], datetime) else None, asof),
                "months_to_expiry": months_between(asof, row[12] if isinstance(row[12], datetime) else None),
            }
            apply(str(row[6] or ""), row[7] if isinstance(row[7], (int, float)) else None)
            continue
        if cur is None:
            continue
        if isinstance(uid, str) and uid.lower().startswith("total"):
            break
        apply(str(row[6] or ""), row[7] if isinstance(row[7], (int, float)) else None)
    flush()
    return [_labeled(u) for u in units if _keep(u)]


def _rediq(path: Path, property_name: str, asof: datetime) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = list(wb["Rent Roll"].iter_rows(values_only=True))
    header = next(i for i, row in enumerate(rows[:15]) if any(str(c or "").strip() == "UnitID" for c in row))
    labels = [str(c or "").strip() for c in rows[header]]
    idx = {name: j for j, name in enumerate(labels) if name}

    def col(*names: str) -> int | None:
        for name in names:
            if name in idx:
                return idx[name]
        return None

    c_unit, c_plan, c_sf = col("UnitID"), col("PlanID"), col("NetSF")
    c_bed, c_occ, c_mkt = col("Bed"), col("OccStatus"), col("MktRent")
    c_ip, c_exp, c_mtm = col("InPlaceRent"), col("LeaseExp"), col("mtm")
    c_move, c_out, c_type = col("MoveIn"), col("MoveOut"), col("UnitType")
    units = []
    for row in rows[header + 1 :]:
        uid = row[c_unit] if c_unit is not None else None
        if not isinstance(uid, str) or not uid.strip() or uid.strip() in {"UnitID", "Unit No."}:
            continue
        occ_raw = str(row[c_occ] or "").strip().lower() if c_occ is not None else ""
        if "occup" in occ_raw:
            occ = "occupied"
        elif "vacant" in occ_raw:
            occ = "vacant"
        elif "non" in occ_raw or "model" in occ_raw:
            occ = "excluded"
        else:
            occ = "other"
        plan = str(row[c_plan] or "").strip() if c_plan is not None else ""
        asset = str(row[c_type] or "").strip().lower() if c_type is not None else ""
        bed = row[c_bed] if c_bed is not None else None
        beds = float(bed) if isinstance(bed, (int, float)) else beds_from_plan(plan)
        move_in = row[c_move] if c_move is not None and isinstance(row[c_move], datetime) else None
        lease_end = row[c_exp] if c_exp is not None and isinstance(row[c_exp], datetime) else None
        move_out = row[c_out] if c_out is not None and isinstance(row[c_out], datetime) else None
        mtm_raw = str(row[c_mtm] or "").strip().upper() if c_mtm is not None else ""
        units.append(
            {
                "property": property_name,
                "as_of": asof.date().isoformat(),
                "plan": plan,
                "sqft": float(row[c_sf]) if c_sf is not None and isinstance(row[c_sf], (int, float)) else None,
                "beds": beds,
                "reno": "reno" in plan.lower(),
                "mtm": mtm_raw in {"MTM", "Y", "YES", "1"},
                "in_place": float(row[c_ip]) if c_ip is not None and isinstance(row[c_ip], (int, float)) else None,
                "market": float(row[c_mkt]) if c_mkt is not None and isinstance(row[c_mkt], (int, float)) else None,
                "occ": occ,
                "notice": move_out is not None,
                "retail": "retail" in asset or "commercial" in asset,
                "tenure_months": months_between(move_in, asof),
                "months_to_expiry": months_between(asof, lease_end),
            }
        )
    wb.close()
    return [_labeled(u) for u in units if _keep(u)]


def _keep(unit: dict) -> bool:
    if unit["occ"] != "occupied":
        return False
    rent = unit.get("in_place")
    sqft = unit.get("sqft")
    if not rent or not sqft:
        return False
    return float(rent) >= 100 and float(sqft) >= 50


def _labeled(unit: dict) -> dict:
    unit["y"] = float(unit["in_place"])
    unit["x"] = feature_row(unit)
    return unit
